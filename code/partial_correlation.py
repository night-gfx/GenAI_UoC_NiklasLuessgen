"""Describe utilities–technology co-movement after controlling for the market.

Task 02 residuals remain separate by period; weekly volatilities retain percent
units (0.02 denotes 0.02 percent). No significance tests are performed here.
"""

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "output"
RESIDUALS_PATH = OUTPUT_DIR / "market_residuals.csv"
STATISTICS_PATH = OUTPUT_DIR / "regression_statistics.csv"
PERIODS = (
    ("Full Sample", "1990-01-01", "2026-08-28"),
    ("Before AI boom", "1990-01-01", "2022-11-30"),
    ("AI boom", "2022-12-01", "2026-08-28"),
    ("Pre-Dot-Com", "1990-01-01", "2000-02-29"),
    ("Dot-com Crash", "2000-03-01", "2002-10-31"),
    ("Pre Financial Crisis", "2002-11-01", "2007-05-31"),
    ("Financial Crisis", "2007-06-01", "2009-02-28"),
    ("Pre-COVID", "2009-03-01", "2020-02-29"),
    ("COVID", "2020-03-01", "2022-02-28"),
)
PERIOD_NAMES = [name for name, _, _ in PERIODS]
RESIDUAL_COLUMNS = [
    "period", "week_end", "utilities_residual", "technology_residual", "market_excess",
]
STATISTICS_REQUIRED = [
    "period", "n_obs", "corr_UT", "corr_UM", "corr_TM", "vol_U", "vol_T", "vol_M",
]
PARTIAL_COLUMNS = ["period", "n_obs", "partial_correlation"]
CHECK_ATOL = 1e-12
CHECK_RTOL = 1e-10
PANEL_A_COLUMNS = [
    "Comparison period", "Partial Correlation", "Corr(U, T)", "Corr(U, M)", "Corr(T, M)", "N",
]
PANEL_B_COLUMNS = ["Comparison period", "σ Utilities", "σ Technology", "σ Market"]
PANEL_TITLES = {
    "Panel A": "Panel A: Correlations",
    "Panel B": "Panel B: Volatilities (standard deviation of weekly excess returns, %)",
}


def _validate_periods(table, name):
    """Require exactly one row per prescribed period, irrespective of input order."""
    if (not table.columns.is_unique or "period" not in table.columns
            or table["period"].isna().any() or table["period"].duplicated().any()
            or set(table["period"]) != set(PERIOD_NAMES)):
        raise ValueError(f"{name}: expected exactly the nine unique proposal periods")


def _validate_residuals(residuals):
    """Validate residual keys/values without dropping, filling or pooling rows."""
    if list(residuals.columns) != RESIDUAL_COLUMNS:
        raise ValueError("Invalid residual columns")
    if (residuals[["period", "week_end"]].isna().any().any()
            or residuals.duplicated(["period", "week_end"]).any()
            or set(residuals["period"]) != set(PERIOD_NAMES)):
        raise ValueError("Missing, duplicate or unknown residual keys")
    if not np.isfinite(residuals[RESIDUAL_COLUMNS[2:]].to_numpy(dtype=float)).all():
        raise ValueError("Nonfinite residual values")


def _validate_statistics(statistics):
    _validate_periods(statistics, "Regression statistics")
    if not set(STATISTICS_REQUIRED).issubset(statistics.columns):
        raise ValueError("Missing regression statistics columns")
    numeric = statistics[STATISTICS_REQUIRED[1:]].to_numpy(dtype=float)
    if not np.isfinite(numeric).all():
        raise ValueError("Nonfinite regression statistics")
    n = statistics["n_obs"].to_numpy(dtype=float)
    if (n < 2).any() or (n != np.floor(n)).any():
        raise ValueError("Invalid regression observation counts")
    if (statistics[["corr_UT", "corr_UM", "corr_TM"]].abs() > 1).any().any():
        raise ValueError("Invalid input correlations")
    if (statistics[["vol_U", "vol_T", "vol_M"]] <= 0).any().any():
        raise ValueError("Nonpositive weekly volatility")


def load_inputs(residuals_path, statistics_path):
    """Read Task 02 inputs; reject missing weeks, invalid keys and count mismatches."""
    residuals = pd.read_csv(residuals_path, float_precision="round_trip")
    statistics = pd.read_csv(statistics_path, float_precision="round_trip")
    _validate_residuals(residuals)
    _validate_statistics(statistics)
    dates = residuals["week_end"].astype("string")
    if not dates.str.fullmatch(r"\d{4}-\d{2}-\d{2}").fillna(False).all():
        raise ValueError("Expected ISO YYYY-MM-DD residual dates")
    residuals["week_end"] = pd.to_datetime(dates, format="%Y-%m-%d", errors="raise")
    keyed_statistics = statistics.set_index("period")
    for period, start, end in PERIODS:
        group = residuals.loc[residuals["period"].eq(period)]
        expected = pd.date_range(start, end, freq="W-FRI")
        if not pd.DatetimeIndex(group["week_end"]).equals(expected):
            raise ValueError(f"{period}: expected every prescribed Friday exactly once in order")
        if len(group) != keyed_statistics.loc[period, "n_obs"]:
            raise ValueError(f"{period}: residual count differs from Task 02 n_obs")
        for column, actual in [("first_week", expected[0]), ("last_week", expected[-1])]:
            if column in statistics and pd.Timestamp(keyed_statistics.loc[period, column]) != actual:
                raise ValueError(f"{period}: {column} differs from residual dates")
    return residuals, statistics


def _validate_partial_correlations(partial_correlations):
    _validate_periods(partial_correlations, "Partial correlations")
    if list(partial_correlations.columns) != PARTIAL_COLUMNS:
        raise ValueError("Invalid partial correlation columns")
    numbers = partial_correlations[PARTIAL_COLUMNS[1:]].to_numpy(dtype=float)
    if not np.isfinite(numbers).all():
        raise ValueError("Nonfinite partial correlations or counts")
    n = partial_correlations["n_obs"].to_numpy(dtype=float)
    if (n < 2).any() or (n != np.floor(n)).any():
        raise ValueError("Invalid partial correlation observation counts")
    if (partial_correlations["partial_correlation"].abs() > 1).any():
        raise ValueError("Partial correlation outside [-1, 1]; no clipping is permitted")


def calculate_partial_correlations(residuals):
    """Compute Pearson correlations within each period, without clipping results."""
    _validate_residuals(residuals)
    rows = []
    for period in PERIOD_NAMES:
        group = residuals.loc[residuals["period"].eq(period)]
        values = group[["utilities_residual", "technology_residual"]].to_numpy(dtype=float)
        # Check the exact range too: constant decimals can have roundoff variance.
        if len(values) < 2 or (np.ptp(values, axis=0) == 0).any():
            raise ValueError(f"{period}: zero residual variance or insufficient observations")
        centered = values - values.mean(axis=0)
        sum_squares = np.sum(centered**2, axis=0)
        if (sum_squares <= 0).any() or not np.isfinite(sum_squares).all():
            raise ValueError(f"{period}: invalid residual variance")
        correlation = np.dot(centered[:, 0], centered[:, 1]) / np.sqrt(sum_squares[0] * sum_squares[1])
        rows.append([period, len(group), correlation])
    partial_correlations = pd.DataFrame(rows, columns=PARTIAL_COLUMNS)
    _validate_partial_correlations(partial_correlations)
    return partial_correlations


def format_descriptive_statistics(partial_correlations, regression_statistics):
    """Join by period, verify the analytic identity, and round only panel displays."""
    _validate_partial_correlations(partial_correlations)
    _validate_statistics(regression_statistics)
    joined = partial_correlations.merge(
        regression_statistics[STATISTICS_REQUIRED], on="period", validate="one_to_one",
        suffixes=("_partial", "_regression"),
    ).set_index("period").loc[PERIOD_NAMES].reset_index()
    if not np.array_equal(joined["n_obs_partial"], joined["n_obs_regression"]):
        raise ValueError("Partial correlation counts differ from Task 02 n_obs")
    denominator_squared = (1 - joined["corr_UM"]**2) * (1 - joined["corr_TM"]**2)
    if (denominator_squared <= 0).any():
        raise ValueError("Analytic partial correlation has a zero or invalid denominator")
    analytic = (
        (joined["corr_UT"] - joined["corr_UM"] * joined["corr_TM"])
        / np.sqrt(denominator_squared)
    )
    matches = np.isclose(joined["partial_correlation"], analytic, atol=CHECK_ATOL, rtol=CHECK_RTOL)
    if not matches.all():
        failed = joined.loc[~matches, "period"].tolist()
        raise ValueError(f"Residual and analytic partial correlations disagree: {failed}")
    panel_a = joined[[
        "period", "partial_correlation", "corr_UT", "corr_UM", "corr_TM", "n_obs_partial",
    ]].copy()
    panel_a.columns = PANEL_A_COLUMNS
    panel_a["N"] = panel_a["N"].astype(int)
    panel_b = joined[["period", "vol_U", "vol_T", "vol_M"]].copy()
    panel_b.columns = PANEL_B_COLUMNS
    return {"Panel A": panel_a.round(3), "Panel B": panel_b.round(3)}


def save_results(panels, partial_correlations, output_dir):
    """Save only the two named outputs, using percent values without rescaling."""
    destination = Path(output_dir).resolve()
    if destination != OUTPUT_DIR:
        raise ValueError(f"The only approved output directory is {OUTPUT_DIR}")
    _validate_partial_correlations(partial_correlations)
    ordered = partial_correlations.set_index("period").loc[PERIOD_NAMES].reset_index()
    if set(panels) != set(PANEL_TITLES):
        raise ValueError("Expected Panel A and Panel B")
    for name, columns in [("Panel A", PANEL_A_COLUMNS), ("Panel B", PANEL_B_COLUMNS)]:
        panel = panels[name]
        if list(panel.columns) != columns or panel.iloc[:, 0].tolist() != PERIOD_NAMES:
            raise ValueError(f"{name}: invalid columns or period order")
        numeric = panel.iloc[:, 1:].to_numpy(dtype=float)
        if not np.isfinite(numeric).all() or not np.array_equal(numeric, np.round(numeric, 3)):
            raise ValueError(f"{name}: expected finite values rounded to three decimals")
    panel_a, panel_b = panels["Panel A"], panels["Panel B"]
    if (not np.array_equal(panel_a["N"], ordered["n_obs"])
            or not np.array_equal(panel_a["Partial Correlation"], ordered["partial_correlation"].round(3))
            or (panel_a.iloc[:, 1:5].abs() > 1).any().any()
            or (panel_b.iloc[:, 1:] < 0).any().any()):
        raise ValueError("Panels contain inconsistent counts, correlations or volatilities")
    excel_path = destination / "descriptive_statistics.xlsx"
    csv_path = destination / "partial_correlations.csv"
    if any(path.resolve() != path for path in (excel_path, csv_path)):
        raise ValueError("Output paths must not redirect to other files")
    destination.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    for name in PANEL_TITLES:
        panel = panels[name]
        sheet = book.create_sheet(name)
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(panel.columns))
        sheet.cell(1, 1, PANEL_TITLES[name])
        sheet.cell(1, 1).font = Font(name="Calibri", size=12, bold=True)
        sheet.cell(1, 1).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 34
        for column, title in enumerate(panel.columns, start=1):
            sheet.cell(3, column, title)
        excel_row = 4
        for i, row in enumerate(panel.itertuples(index=False, name=None)):
            if i in (1, 3):
                excel_row += 1
            for column, value in enumerate(row, start=1):
                sheet.cell(excel_row, column, value)
            excel_row += 1
        for row in sheet.iter_rows(min_row=3, max_row=excel_row - 1):
            for cell in row:
                cell.font = Font(name="Calibri", size=11, bold=cell.row == 3)
                cell.alignment = Alignment(horizontal="left" if cell.column == 1 else "right")
                if cell.row in (3, excel_row - 1):
                    cell.border = Border(bottom=Side(style="thin", color="404040"))
                if cell.row > 3 and cell.column > 1:
                    cell.number_format = "0" if name == "Panel A" and cell.column == 6 else "0.000"
        widths = [26, 22, 17, 17, 17, 9] if name == "Panel A" else [26, 20, 20, 20]
        for column, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.freeze_panes = "B4"
        sheet.sheet_view.showGridLines = False
        sheet.print_area = f"A1:{get_column_letter(len(panel.columns))}{excel_row - 1}"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
    book.save(excel_path)
    ordered.to_csv(csv_path, index=False)


def main():
    """Load → partial correlation → analytic validation/format → save."""
    residuals, statistics = load_inputs(RESIDUALS_PATH, STATISTICS_PATH)
    partial_correlations = calculate_partial_correlations(residuals)
    panels = format_descriptive_statistics(partial_correlations, statistics)
    save_results(panels, partial_correlations, OUTPUT_DIR)
    print(f"Saved {len(partial_correlations)} periods to {OUTPUT_DIR}")
    print("descriptive_statistics.xlsx, partial_correlations.csv")


if __name__ == "__main__":
    main()
