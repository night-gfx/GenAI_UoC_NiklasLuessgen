"""Estimate period-specific market betas from weekly excess returns.

Returns remain in source percent units: 0.02 denotes 0.02 percent.
"""

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from scipy.stats import norm
from statsmodels.regression.linear_model import OLS


PROJECT_ROOT = Path(__file__).resolve().parents[1]
WEEKLY_PATH = PROJECT_ROOT / "output/weekly_returns.csv"
OUTPUT_DIR = PROJECT_ROOT / "output"
SAMPLE_START = pd.Timestamp("1990-01-01")
SAMPLE_END = pd.Timestamp("2026-08-28")
HAC_LAGS = 4
RETURN_COLUMNS = ["utilities_excess", "technology_excess", "market_excess"]
PERIODS = (
    ("Full Sample", "1990-01", "2026-08"),
    ("Before AI boom", "1990-01", "2022-11"),
    ("AI boom", "2022-12", "2026-08"),
    ("Pre-Dot-Com", "1990-01", "2000-02"),
    ("Dot-com Crash", "2000-03", "2002-10"),
    ("Pre Financial Crisis", "2002-11", "2007-05"),
    ("Financial Crisis", "2007-06", "2009-02"),
    ("Pre-COVID", "2009-03", "2020-02"),
    ("COVID", "2020-03", "2022-02"),
)
PERIOD_NAMES = [name for name, _, _ in PERIODS]
SECTOR_STATISTICS = [
    "alpha", "se_alpha", "beta", "se_beta", "t_beta_1", "p_beta_1",
    "r_squared", "residual_std",
]
STATISTICS_COLUMNS = [
    "period", "n_obs", "first_week", "last_week", "hac_lags",
    *[f"{stat}_{sector}" for sector in ("U", "T") for stat in SECTOR_STATISTICS],
    "vol_U", "vol_T", "vol_M", "corr_UM", "corr_TM", "corr_UT",
]
RESIDUAL_COLUMNS = [
    "period", "week_end", "utilities_residual", "technology_residual", "market_excess",
]
TABLE_COLUMNS = [
    "Comparison period", "Utilities Beta (SE)", "Technology Beta (SE)",
    "R² Utilities", "R² Technology", "N",
]
TABLE_NOTE = (
    "Newey–West standard errors in parentheses; stars test H0: β = 1 (two-sided): "
    "* p < 0.10, ** p < 0.05, *** p < 0.01."
)


def load_weekly_returns(weekly_path):
    """Validate all fixed-sample Fridays and key unchanged returns by period/date.

    The six crisis subperiods partition Before AI boom except March–November
    2022. Full Sample and the before/AI split intentionally repeat observations.
    """
    weekly = pd.read_csv(weekly_path, float_precision="round_trip")
    if list(weekly.columns) != ["week_end", *RETURN_COLUMNS]:
        raise ValueError("Expected week_end and the three weekly excess-return columns")
    dates = weekly["week_end"].astype("string")
    if not dates.str.fullmatch(r"\d{4}-\d{2}-\d{2}").fillna(False).all():
        raise ValueError("Expected ISO YYYY-MM-DD week_end dates")
    weekly["week_end"] = pd.to_datetime(dates, format="%Y-%m-%d", errors="raise")
    expected = pd.date_range(SAMPLE_START, SAMPLE_END, freq="W-FRI")
    if not pd.DatetimeIndex(weekly["week_end"]).equals(expected):
        raise ValueError("Expected every fixed-sample Friday once, in chronological order")
    weekly[RETURN_COLUMNS] = weekly[RETURN_COLUMNS].apply(pd.to_numeric, errors="raise")
    if not np.isfinite(weekly[RETURN_COLUMNS].to_numpy(dtype=float)).all():
        raise ValueError("Nonfinite weekly excess returns; no rows may be filled or dropped")
    weekly = weekly.set_index("week_end")
    months = weekly.index.to_period("M")
    masks = {
        name: (months >= pd.Period(start, "M")) & (months <= pd.Period(end, "M"))
        for name, start, end in PERIODS
    }
    if any(not mask.any() for mask in masks.values()):
        raise ValueError("Empty comparison period")
    split = masks["Before AI boom"].astype(int) + masks["AI boom"].astype(int)
    crisis_count = np.sum([masks[name] for name in PERIOD_NAMES[3:]], axis=0)
    gap = (months >= pd.Period("2022-03", "M")) & (months <= pd.Period("2022-11", "M"))
    if not masks["Full Sample"].all() or not (split == 1).all():
        raise ValueError("Invalid Full Sample or before/AI partition")
    if not np.array_equal(crisis_count + gap, masks["Before AI boom"].astype(int)):
        raise ValueError("Overlapping or incomplete crisis subperiods")
    periods = pd.concat(
        [weekly.loc[masks[name]] for name in PERIOD_NAMES],
        keys=PERIOD_NAMES, names=["period", "week_end"],
    )
    periods.attrs["units"] = "percent"
    return periods


def run_ols(weekly_periods):
    """Refit both sectors per period using OLS and four-lag Newey–West SEs.

    HAC uses Bartlett weights and no finite-sample multiplier (the asymptotic
    Newey–West estimator). Tests use the standard normal distribution.
    residual_std is the sample standard deviation of residuals (ddof=1),
    like the three reported weekly volatilities, rather than a regression RMSE.
    """
    if (not isinstance(weekly_periods.index, pd.MultiIndex)
            or weekly_periods.index.names != ["period", "week_end"]
            or not weekly_periods.index.is_unique
            or list(weekly_periods.columns) != RETURN_COLUMNS
            or list(weekly_periods.index.get_level_values("period").unique()) != PERIOD_NAMES):
        raise ValueError("Expected unique (period, week_end) keys in proposal order")
    if not np.isfinite(weekly_periods.to_numpy(dtype=float)).all():
        raise ValueError("Nonfinite regression input")
    rows, residual_frames = [], []
    for period in PERIOD_NAMES:
        weekly = weekly_periods.loc[period]
        dates = weekly.index
        if (not isinstance(dates, pd.DatetimeIndex) or dates.hasnans
                or not dates.is_monotonic_increasing or (dates.dayofweek != 4).any()
                or len(weekly) <= HAC_LAGS):
            raise ValueError(f"{period}: invalid weekly dates or insufficient observations")
        values = weekly.to_numpy(dtype=float)
        # A constant decimal can have a tiny positive variance from rounding.
        if (np.ptp(values, axis=0) == 0).any() or (np.var(values, axis=0, ddof=1) <= 0).any():
            raise ValueError(f"{period}: zero return variance")
        market = values[:, 2]
        design = np.column_stack([np.ones(len(weekly)), market])
        if np.linalg.matrix_rank(design) != 2:
            raise ValueError(f"{period}: singular regression design")
        row = {
            "period": period, "n_obs": len(weekly), "first_week": dates[0],
            "last_week": dates[-1], "hac_lags": HAC_LAGS,
        }
        residual_frame = pd.DataFrame({"period": period, "week_end": dates})
        for i, (sector, residual_column) in enumerate([
            ("U", "utilities_residual"), ("T", "technology_residual"),
        ]):
            fit = OLS(values[:, i], design, missing="raise").fit(
                cov_type="HAC", use_t=False,
                cov_kwds={"maxlags": HAC_LAGS, "kernel": "bartlett", "use_correction": False},
            )
            alpha, beta = fit.params
            se_alpha, se_beta = fit.bse
            if se_beta <= 0:
                raise ValueError(f"{period}/{sector}: nonpositive beta standard error")
            t_beta_1 = (beta - 1) / se_beta
            stats = [
                alpha, se_alpha, beta, se_beta, t_beta_1, 2 * norm.sf(abs(t_beta_1)),
                fit.rsquared, np.std(fit.resid, ddof=1),
            ]
            row.update({f"{name}_{sector}": value for name, value in zip(SECTOR_STATISTICS, stats)})
            residual_frame[residual_column] = fit.resid
        row.update(dict(zip(["vol_U", "vol_T", "vol_M"], np.std(values, axis=0, ddof=1))))
        correlations = np.corrcoef(values, rowvar=False)
        row.update(corr_UM=correlations[0, 2], corr_TM=correlations[1, 2], corr_UT=correlations[0, 1])
        residual_frame["market_excess"] = market
        rows.append(row)
        residual_frames.append(residual_frame)
    statistics = pd.DataFrame(rows, columns=STATISTICS_COLUMNS)
    residuals = pd.concat(residual_frames, ignore_index=True)[RESIDUAL_COLUMNS]
    numeric = statistics.drop(columns=["period", "first_week", "last_week"])
    if (not np.isfinite(numeric.to_numpy(dtype=float)).all()
            or not np.isfinite(residuals[RESIDUAL_COLUMNS[2:]].to_numpy(dtype=float)).all()):
        raise ValueError("Nonfinite regression results")
    return statistics, residuals


def _validate_statistics(statistics):
    """Reject missing, unordered or invalid statistics before table/export use."""
    if (list(statistics.columns) != STATISTICS_COLUMNS
            or statistics["period"].tolist() != PERIOD_NAMES):
        raise ValueError("Expected the complete statistics schema and nine ordered periods")
    numeric = statistics.drop(columns=["period", "first_week", "last_week"])
    if not np.isfinite(numeric.to_numpy(dtype=float)).all():
        raise ValueError("Nonfinite regression statistics")
    n = statistics["n_obs"]
    if (n <= HAC_LAGS).any() or (n != np.floor(n)).any() or not statistics["hac_lags"].eq(HAC_LAGS).all():
        raise ValueError("Invalid observation counts or HAC lag length")
    first = pd.to_datetime(statistics["first_week"], errors="raise")
    last = pd.to_datetime(statistics["last_week"], errors="raise")
    if first.isna().any() or last.isna().any() or (first > last).any():
        raise ValueError("Invalid first/last weeks")
    for sector in ("U", "T"):
        if ((statistics[f"se_beta_{sector}"] <= 0).any()
                or (statistics[f"se_alpha_{sector}"] < 0).any()
                or (statistics[f"residual_std_{sector}"] < 0).any()
                or not statistics[f"p_beta_1_{sector}"].between(0, 1).all()
                or not statistics[f"r_squared_{sector}"].between(-1e-12, 1 + 1e-12).all()):
            raise ValueError(f"Invalid {sector} standard errors, p-values or R²")
    if (statistics[["vol_U", "vol_T", "vol_M"]] <= 0).any().any():
        raise ValueError("Nonpositive volatility")
    if (statistics[["corr_UM", "corr_TM", "corr_UT"]].abs() > 1 + 1e-12).any().any():
        raise ValueError("Invalid Pearson correlations")


def format_regression_table(regression_statistics):
    """Return nine proposal rows, rounding only the displayed table to 3 decimals."""
    _validate_statistics(regression_statistics)
    rows = []
    for row in regression_statistics.itertuples(index=False):
        betas = []
        for sector in ("U", "T"):
            p = getattr(row, f"p_beta_1_{sector}")
            stars = "***" if p < 0.01 else "**" if p < 0.05 else "*" if p < 0.10 else ""
            betas.append(f"{getattr(row, f'beta_{sector}'):.3f}{stars} ({getattr(row, f'se_beta_{sector}'):.3f})")
        rows.append([
            row.period, *betas, round(row.r_squared_U, 3), round(row.r_squared_T, 3), int(row.n_obs),
        ])
    table = pd.DataFrame(rows, columns=TABLE_COLUMNS)
    table.attrs["note"] = TABLE_NOTE
    return table


def save_results(regression_table, regression_statistics, residuals, output_dir):
    """Save only the three authorized outputs; propagate validation/write errors."""
    destination = Path(output_dir).resolve()
    if destination != OUTPUT_DIR:
        raise ValueError(f"The only approved output directory is {OUTPUT_DIR}")
    expected_table = format_regression_table(regression_statistics)
    if (not regression_table.equals(expected_table)
            or regression_table.attrs.get("note") != TABLE_NOTE):
        raise ValueError("Regression table does not match the unrounded statistics")
    if (list(residuals.columns) != RESIDUAL_COLUMNS
            or residuals.duplicated(["period", "week_end"]).any()
            or residuals[["period", "week_end"]].isna().any().any()
            or not np.isfinite(residuals[RESIDUAL_COLUMNS[2:]].to_numpy(dtype=float)).all()):
        raise ValueError("Invalid residual schema, keys or values")
    counts = residuals.groupby("period", sort=False).size()
    if (counts.index.tolist() != PERIOD_NAMES
            or not np.array_equal(counts.to_numpy(), regression_statistics["n_obs"].to_numpy())):
        raise ValueError("Residual counts differ from regression statistics")
    dates = pd.to_datetime(residuals["week_end"], errors="raise")
    if (dates.dt.dayofweek != 4).any():
        raise ValueError("Residual dates must be Fridays")
    paths = [destination / name for name in (
        "regression_table.xlsx", "regression_statistics.csv", "market_residuals.csv",
    )]
    if any(path.resolve() != path for path in paths):
        raise ValueError("Output paths must not redirect to other files")
    destination.mkdir(parents=True, exist_ok=True)

    book = Workbook()
    sheet = book.active
    sheet.title = "Market Betas"
    sheet.append(TABLE_COLUMNS)
    for i, row in enumerate(regression_table.itertuples(index=False, name=None)):
        if i in (1, 3):
            sheet.append([None] * len(TABLE_COLUMNS))
        sheet.append(row)
    last_data_row = sheet.max_row
    note_row = last_data_row + 2
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=6)
    sheet.cell(note_row, 1, TABLE_NOTE)
    sheet.cell(note_row, 1).font = Font(name="Calibri", size=10, italic=True)
    sheet.cell(note_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[note_row].height = 32
    for row in sheet.iter_rows(min_row=1, max_row=last_data_row):
        for cell in row:
            cell.font = Font(name="Calibri", size=11, bold=cell.row == 1)
            cell.alignment = Alignment(horizontal="left" if cell.column == 1 else "right")
            if cell.row == 1 or cell.row == last_data_row:
                cell.border = Border(bottom=Side(style="thin", color="404040"))
            if cell.row > 1 and cell.column in (4, 5):
                cell.number_format = "0.000"
            elif cell.row > 1 and cell.column == 6:
                cell.number_format = "0"
    for i, width in enumerate([26, 25, 25, 17, 17, 9], start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.freeze_panes = "B2"
    sheet.sheet_view.showGridLines = False
    sheet.print_area = f"A1:F{note_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    book.save(paths[0])
    regression_statistics.to_csv(paths[1], index=False, date_format="%Y-%m-%d")
    residuals.to_csv(paths[2], index=False, date_format="%Y-%m-%d")


def main():
    """Load → OLS → format → save, with no input scaling or row exclusions."""
    weekly_periods = load_weekly_returns(WEEKLY_PATH)
    statistics, residuals = run_ols(weekly_periods)
    table = format_regression_table(statistics)
    save_results(table, statistics, residuals, OUTPUT_DIR)
    print(f"Saved {len(statistics)} periods and {len(residuals)} residual rows to {OUTPUT_DIR}")
    print("regression_table.xlsx, regression_statistics.csv, market_residuals.csv")


if __name__ == "__main__":
    main()
